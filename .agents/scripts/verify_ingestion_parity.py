"""
verify_ingestion_parity.py — Lossless Data Ingestion Parity Verifier (Rule 27)
=============================================================================
Audits whether empirical data (Sample Size N, Demographics, Item Counts,
Likert Anchors, Means, SDs, Reliabilities, and Correlations r) extracted by
Subagents into structured JSON were losslessly transcribed into the target
Excel coding sheet.

If any valid empirical metric in the JSON is recorded as '999' (or dropped),
if any numeric value mismatches between JSON and Excel, or if the inclusion/exclusion
verdict mismatches, this script flags a CRITICAL_PARITY_VIOLATION and exits with code 1.
"""

import os
import re
import sys
import json
import argparse
from typing import Dict, Any, List, Tuple, Optional
import openpyxl

# Force UTF-8 output
sys.stdout.reconfigure(encoding='utf-8')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
if SCRIPT_DIR not in sys.path:
    sys.path.append(SCRIPT_DIR)

from universal_excel_inserter import (
    sanitize_numeric,
    smart_anchor_match,
    extract_screening_verdict,
    normalize_measure_details,
)


def normalize_anchor(s: str) -> str:
    """Normalize variable name for comparison (pruning leading layout indices per Rule 14)."""
    if not s:
        return ""
    s_clean = re.sub(r'^\s*(?:\(\d+\)|\d+[\.\)])\s*', '', str(s).strip().lower())
    s_clean = s_clean.replace(".", "").replace("(", "").replace(")", "").replace(" ", "").replace("_", "").replace("-", "")
    return s_clean


def parse_metric(v: Any) -> Optional[float]:
    """Parse value using sanitize_numeric; return float if valid (not 999), else None."""
    res = sanitize_numeric(v, default=999)
    if res == 999 or res is None or isinstance(res, bool):
        return None
    try:
        f_val = float(res)
        return f_val if abs(f_val - 999.0) > 1e-4 else None
    except (ValueError, TypeError):
        return None


def is_valid_metric(v: Any) -> bool:
    """Check if value is a valid numeric metric (not None, not empty, not 999)."""
    return parse_metric(v) is not None


def extract_reliability_val(obj: Any) -> Any:
    """Extract numeric reliability value from dict or scalar."""
    if isinstance(obj, dict):
        return obj.get("value")
    return obj


def find_var_entry(anchor_name: str, json_vars: Dict[str, Dict[str, Any]]) -> Optional[Dict[str, Any]]:
    """Find matching variable entry in json_vars by exact normalized key or smart_anchor_match."""
    norm = normalize_anchor(anchor_name)
    if norm in json_vars:
        return json_vars[norm]
    for k, entry in json_vars.items():
        if smart_anchor_match(anchor_name, entry.get("raw_name", "")):
            return entry
    return None


def find_corr_entries(bs_name: str, nb_name: str, json_corrs: List[Tuple[str, str, float]]) -> List[float]:
    """Find all matching correlation r values between bs_name and nb_name across studies."""
    bs_norm = normalize_anchor(bs_name)
    nb_norm = normalize_anchor(nb_name)
    exact_matches = [
        r_num for v1_raw, v2_raw, r_num in json_corrs
        if (normalize_anchor(v1_raw) == bs_norm and normalize_anchor(v2_raw) == nb_norm) or
           (normalize_anchor(v1_raw) == nb_norm and normalize_anchor(v2_raw) == bs_norm)
    ]
    if exact_matches:
        return exact_matches
    return [
        r_num for v1_raw, v2_raw, r_num in json_corrs
        if (smart_anchor_match(bs_name, v1_raw) and smart_anchor_match(nb_name, v2_raw)) or
           (smart_anchor_match(bs_name, v2_raw) and smart_anchor_match(nb_name, v1_raw))
    ]


def verify_paper_parity(excel_path: str, json_path: str, paper_id: int) -> Tuple[bool, List[str], List[str]]:
    """
    Compare JSON extraction output against Excel sheet data for a specific paper.
    Returns: (passed, violations, confirmations)
    """
    violations = []
    confirmations = []

    if not os.path.exists(excel_path):
        return False, [f"Excel file not found: {excel_path}"], []
    if not os.path.exists(json_path):
        return False, [f"JSON file not found: {json_path}"], []

    with open(json_path, 'r', encoding='utf-8') as f:
        payload = json.load(f)

    # 1. Load Excel sheet and locate rows for paper_id
    wb = openpyxl.load_workbook(excel_path, data_only=True)
    ws = wb["Extracted Data"] if "Extracted Data" in wb.sheetnames else wb.active

    paper_rows = []
    for r in range(4, ws.max_row + 1):
        cell_id = ws.cell(r, 2).value  # Col 2: Article ID
        if cell_id is not None:
            m = re.search(r'(\d+)', str(cell_id))
            if m and int(m.group(1)) == int(paper_id):
                paper_rows.append(r)

    if not paper_rows:
        wb.close()
        return False, [f"No rows found in Excel sheet for Paper ID {paper_id}"], []

    first_row = paper_rows[0]
    judgment_val = str(ws.cell(first_row, 5).value or "").strip()
    excel_is_excluded = judgment_val.startswith("0") or judgment_val.lower().startswith("exclude")
    json_is_excluded, _ = extract_screening_verdict(payload)

    if json_is_excluded != excel_is_excluded:
        wb.close()
        return False, [
            f"Inclusion/Exclusion Verdict mismatch for Paper [{paper_id}]: "
            f"JSON is {'0 = exclude' if json_is_excluded else '1 = include'}, "
            f"but Excel Row {first_row} Col 5 has '{judgment_val}'."
        ], []

    if excel_is_excluded:
        # Also verify Cols 7-50 are blank per Rule 19
        non_blank = [c for c in range(7, 51) if ws.cell(first_row, c).value not in (None, "")]
        if non_blank:
            violations.append(f"Excluded Paper [{paper_id}] Row {first_row} has non-blank cells in Cols 7-50: {non_blank}")
        else:
            confirmations.append(f"Paper [{paper_id}] is Excluded ('{judgment_val}'); Col 7-50 blank check verified.")
        wb.close()
        return len(violations) == 0, violations, confirmations

    # 2. Unpack study payloads (supporting single-study or multi-study 'studies' list)
    studies_list = payload.get("studies")
    if isinstance(studies_list, list) and studies_list:
        study_payloads = []
        for st in studies_list:
            if isinstance(st, dict):
                merged = dict(payload)
                merged.pop("studies", None)
                merged.update(st)
                study_payloads.append(merged)
    else:
        study_payloads = [payload]

    json_vars: Dict[str, Dict[str, Any]] = {}
    json_corrs: List[Tuple[str, str, float]] = []
    expected_ns: List[float] = []
    expected_ages: List[float] = []
    expected_females: List[float] = []
    expected_tenures: List[float] = []

    def merge_var_entry(name: str, data: Dict[str, Any]):
        if not name:
            return
        norm = normalize_anchor(name)
        if not norm:
            return
        entry = json_vars.setdefault(norm, {"raw_name": name})
        for k, v in data.items():
            raw_v = extract_reliability_val(v) if k == "reliability" else v
            m_val = parse_metric(raw_v)
            if m_val is not None:
                vals_list = entry.setdefault(f"{k}_vals", [])
                if not any(abs(m_val - existing) <= 1e-4 for existing in vals_list):
                    vals_list.append(m_val)
                entry[k] = v
            elif k not in entry:
                entry[k] = v

    for st_payload in study_payloads:
        matrix_block = st_payload.get("table_matrix") or st_payload.get("boundary_spanning_matrix") or {}
        if not isinstance(matrix_block, dict):
            matrix_block = {}

        vars_data = st_payload.get("variables") if st_payload.get("variables") is not None else matrix_block.get("variables")
        if isinstance(vars_data, dict):
            for k, v in vars_data.items():
                if isinstance(v, dict):
                    name = v.get("name") or v.get("table_anchor_name") or str(k)
                    merge_var_entry(name, v)
        elif isinstance(vars_data, list):
            for v in vars_data:
                if isinstance(v, dict):
                    name = v.get("table_anchor_name") or v.get("name") or ""
                    merge_var_entry(name, v)

        for m in normalize_measure_details(st_payload):
            name = m.get("table_anchor_name") or m.get("name") or ""
            merge_var_entry(name, m)

        corrs_data = st_payload.get("correlations") if st_payload.get("correlations") is not None else matrix_block.get("correlations")
        if isinstance(corrs_data, list):
            for c in corrs_data:
                if isinstance(c, dict):
                    v1 = c.get("var1_anchor", "")
                    v2 = c.get("var2_anchor", "")
                    r_num = parse_metric(c.get("r"))
                    if v1 and v2 and r_num is not None:
                        json_corrs.append((v1, v2, r_num))

        sample_blocks = [
            st_payload.get("study_sample") or {},
            st_payload.get("sample_descriptors") or {},
            st_payload.get("study_sample_descriptor") or {},
            st_payload.get("study_details") or {},
            st_payload.get("design") or {},
            st_payload.get("node_0_screening") or {},
            st_payload.get("screening") or {},
            st_payload,
        ]

        def get_st_metric(*keys) -> Optional[float]:
            for blk in sample_blocks:
                if not isinstance(blk, dict):
                    continue
                for k in keys:
                    if k in blk:
                        m_val = parse_metric(blk[k])
                        if m_val is not None:
                            return m_val
            return None

        n_val = get_st_metric("listwise_n", "sample_size_n", "sample_n", "effective_sample_size", "effective_n")
        if n_val is not None:
            expected_ns.append(n_val)
        age_val = get_st_metric("mean_age")
        if age_val is not None:
            expected_ages.append(age_val)
        fem_val = get_st_metric("pct_female", "percent_female")
        if fem_val is not None:
            expected_females.append(fem_val)
        ten_val = get_st_metric("org_tenure", "tenure")
        if ten_val is not None:
            expected_tenures.append(ten_val)

    # 3. Verify Sample Size N (Col 22) and Demographics (Cols 23, 24, 25) across paper_rows
    for label, col_idx, exp_list in [
        ("Sample Size N (Col 22)", 22, expected_ns),
        ("Demographics Mean Age (Col 23)", 23, expected_ages),
        ("Demographics % Female (Col 24)", 24, expected_females),
        ("Demographics Org Tenure (Col 25)", 25, expected_tenures),
    ]:
        if exp_list:
            excel_vals_raw = [ws.cell(r, col_idx).value for r in paper_rows]
            excel_nums = [parse_metric(v) for v in excel_vals_raw if parse_metric(v) is not None]
            if not excel_nums:
                violations.append(f"{label}: JSON has {exp_list}, but Excel has '{excel_vals_raw[0]}' (dropped to 999).")
            else:
                for exp_v in exp_list:
                    if not any(abs(exp_v - ex_n) <= 1e-3 for ex_n in excel_nums):
                        violations.append(f"{label}: JSON has '{exp_v}', but Excel has {sorted(set(excel_nums))} (value mismatch).")
                    else:
                        confirmations.append(f"{label} verified: {exp_v}")

    # 4. Verify Variable Measure Descriptors, Means, SDs, Reliabilities, and Correlations across rows
    def get_var_candidates(j_var: Dict[str, Any], *keys: str) -> List[float]:
        out: List[float] = []
        for k in keys:
            for v in j_var.get(f"{k}_vals", []):
                if not any(abs(v - x) <= 1e-4 for x in out):
                    out.append(v)
        return out

    def check_cell_parity(row_num: int, prefix: str, field_lbl: str, col_i: int, candidate_nums: List[float]):
        if not candidate_nums:
            return
        ex_raw = ws.cell(row_num, col_i).value
        ex_num = parse_metric(ex_raw)
        if ex_num is None:
            violations.append(f"Row {row_num} {prefix} {field_lbl}: JSON has {candidate_nums}, but Excel has '{ex_raw}' (dropped to 999).")
        elif not any(abs(c_num - ex_num) <= 1e-3 for c_num in candidate_nums):
            violations.append(f"Row {row_num} {prefix} {field_lbl}: JSON has {candidate_nums}, but Excel has '{ex_num}' (value mismatch).")
        else:
            confirmations.append(f"Row {row_num} {prefix} {field_lbl} verified: {ex_raw}")

    for r in paper_rows:
        bs_name = ws.cell(r, 41).value
        nb_name = ws.cell(r, 45).value

        # Check BSB variable (Cols 27, 28, 29, 42, 43, 44)
        if bs_name:
            j_var = find_var_entry(str(bs_name), json_vars)
            if j_var:
                for field_lbl, col_i, cands in [
                    ("Items (Col 27)", 27, get_var_candidates(j_var, "number_of_items", "items")),
                    ("Min Score (Col 28)", 28, get_var_candidates(j_var, "min_score", "min")),
                    ("Max Score (Col 29)", 29, get_var_candidates(j_var, "max_score", "max")),
                    ("Mean (Col 42)", 42, get_var_candidates(j_var, "mean")),
                    ("SD (Col 43)", 43, get_var_candidates(j_var, "sd")),
                    ("Reliability (Col 44)", 44, get_var_candidates(j_var, "reliability", "reliability_table", "alpha")),
                ]:
                    check_cell_parity(r, f"BSB '{bs_name}'", field_lbl, col_i, cands)

        # Check Non-BS variable (Cols 34, 35, 36, 46, 47, 48)
        if nb_name:
            j_var = find_var_entry(str(nb_name), json_vars)
            if j_var:
                for field_lbl, col_i, cands in [
                    ("Items (Col 34)", 34, get_var_candidates(j_var, "number_of_items", "items")),
                    ("Min Score (Col 35)", 35, get_var_candidates(j_var, "min_score", "min")),
                    ("Max Score (Col 36)", 36, get_var_candidates(j_var, "max_score", "max")),
                    ("Mean (Col 46)", 46, get_var_candidates(j_var, "mean")),
                    ("SD (Col 47)", 47, get_var_candidates(j_var, "sd")),
                    ("Reliability (Col 48)", 48, get_var_candidates(j_var, "reliability", "reliability_table", "alpha")),
                ]:
                    check_cell_parity(r, f"Non-BS '{nb_name}'", field_lbl, col_i, cands)

        # Check Correlation r (Col 49)
        if bs_name and nb_name:
            exp_rs = find_corr_entries(str(bs_name), str(nb_name), json_corrs)
            if exp_rs:
                check_cell_parity(r, f"Pair ('{bs_name}' x '{nb_name}')", "Correlation r (Col 49)", 49, exp_rs)

    wb.close()

    if not confirmations and not violations:
        violations.append(f"Zero metrics confirmed for Included Paper [{paper_id}] — schema or variable mismatch between JSON and Excel.")

    passed = len(violations) == 0
    return passed, violations, confirmations


def main():
    parser = argparse.ArgumentParser(description="Lossless Ingestion Parity Verifier (Rule 27)")
    parser.add_argument("--excel", required=True, help="Path to Excel coding sheet")
    parser.add_argument("--json", required=True, help="Path to Subagent extraction JSON")
    parser.add_argument("--paper-id", type=int, required=True, help="Article ID of the paper")

    args = parser.parse_args()

    print("=" * 70)
    print(f"BSMA Ingestion Parity Audit — Paper ID: [{args.paper_id}]")
    print(f"Excel Sheet: {args.excel}")
    print(f"Source JSON: {args.json}")
    print("=" * 70)

    passed, violations, confirmations = verify_paper_parity(args.excel, args.json, args.paper_id)

    if confirmations:
        print(f"\n[CONFIRMED CHECKS] ({len(confirmations)} verified)")
        for c in confirmations[:12]:
            print(f"  ✓ {c}")

    if not passed:
        print(f"\n[CRITICAL PARITY VIOLATIONS] ({len(violations)} data loss/mismatch incidents detected):")
        for v in violations:
            print(f"  ✗ {v}")
        print("\n[RESULT] FAILED — Data was lost, mismatched, or converted to 999 during ingestion. Commit blocked.")
        sys.exit(1)
    else:
        print("\n[RESULT] PASSED — Zero data drop detected. All JSON empirical metrics match Excel cells perfectly.")
        sys.exit(0)


if __name__ == "__main__":
    main()
