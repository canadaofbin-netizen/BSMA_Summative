"""
excel_template_util.py — Canonical 3-Tier Excel Coding Sheet Generator & Validator
================================================================================
Ensures all generated coding sheets strictly inherit the canonical 3-tier header
structure from 03_Coding_Sheets/49_53_66.xlsx:
- Row 1: High-Level Domain/Section Category (Article, Study/Sample, Measures, Effect Size)
- Row 2: Sub-category Headers (Boundary Spanning, Non-BS, Correlation r, Notes)
- Row 3: Leaf Column Names (Col 1 to Col 50)
- Row 4+: Data rows strictly begin here.

Strictly prohibits flat pandas exports (df.to_excel) which produce 'Unnamed' headers.
"""

import os
import sys
import copy
import argparse
import openpyxl
from openpyxl.utils import range_boundaries, get_column_letter

# Canonical template source
DEFAULT_TEMPLATE_PATH = os.path.join(
    os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
    "03_Coding_Sheets",
    "49_53_66.xlsx"
)

def copy_canonical_headers(src_ws, dest_ws):
    """
    Copies the first 3 header rows, cell formatting, merged ranges,
    and column dimensions from src_ws to dest_ws.
    """
    # 1. Copy values and formatting for Rows 1-3
    for r in range(1, 4):
        dest_ws.row_dimensions[r].height = src_ws.row_dimensions[r].height
        for c in range(1, 51):
            src_cell = src_ws.cell(row=r, column=c)
            dest_cell = dest_ws.cell(row=r, column=c, value=src_cell.value)
            if src_cell.has_style:
                dest_cell.font = copy.copy(src_cell.font)
                dest_cell.border = copy.copy(src_cell.border)
                dest_cell.fill = copy.copy(src_cell.fill)
                dest_cell.number_format = copy.copy(src_cell.number_format)
                dest_cell.protection = copy.copy(src_cell.protection)
                dest_cell.alignment = copy.copy(src_cell.alignment)

    # 2. Replicate merged ranges within Rows 1-3
    for range_ in src_ws.merged_cells.ranges:
        min_col, min_row, max_col, max_row = range_boundaries(str(range_))
        if max_row <= 3:
            dest_ws.merge_cells(
                start_row=min_row,
                start_column=min_col,
                end_row=max_row,
                end_column=max_col
            )

    # 3. Replicate column widths across all 50 columns
    for c in range(1, 51):
        col_letter = get_column_letter(c)
        if col_letter in src_ws.column_dimensions:
            dest_ws.column_dimensions[col_letter].width = src_ws.column_dimensions[col_letter].width

def create_blank_coding_sheet(dest_path, template_path=None, sheet_title="Extracted Data"):
    """
    Creates a new workbook at dest_path with the canonical 3-tier header template.
    Returns the created openpyxl Workbook object.
    """
    if template_path is None:
        template_path = DEFAULT_TEMPLATE_PATH

    if not os.path.exists(template_path):
        raise FileNotFoundError(f"Canonical template not found at '{template_path}'")

    src_wb = openpyxl.load_workbook(template_path)
    src_ws = src_wb.active

    dest_wb = openpyxl.Workbook()
    dest_ws = dest_wb.active
    dest_ws.title = sheet_title

    copy_canonical_headers(src_ws, dest_ws)
    dest_wb.save(dest_path)
    return dest_wb

def validate_coding_sheet_headers(file_path):
    """
    Validates that file_path conforms to the 3-tier header protocol:
    - Exactly 50 columns
    - No 'Unnamed' values in rows 1-3
    - Correct key header anchors
    Returns (is_valid: bool, issues: list)
    """
    if not os.path.exists(file_path):
        return False, [f"File '{file_path}' does not exist."]

    try:
        wb = openpyxl.load_workbook(file_path, data_only=True)
        ws = wb.active
    except Exception as e:
        return False, [f"Failed to load workbook: {e}"]

    issues = []
    if ws.max_column != 50:
        issues.append(f"Column count is {ws.max_column}, expected exactly 50.")

    # Check for 'Unnamed' in rows 1 to 3
    for r in range(1, 4):
        for c in range(1, min(ws.max_column + 1, 51)):
            val = str(ws.cell(r, c).value or "")
            if "unnamed" in val.lower():
                issues.append(f"Row {r}, Col {c} contains 'Unnamed' header: '{val}'")

    # Check key anchors
    r1_col1 = ws.cell(1, 1).value
    r1_col7 = ws.cell(1, 7).value
    r3_col8 = ws.cell(3, 8).value

    if r1_col1 != "Coder Initials":
        issues.append(f"Row 1, Col 1 anchor mismatch. Expected 'Coder Initials', got '{r1_col1}'")
    if r1_col7 != "Article Descriptors":
        issues.append(f"Row 1, Col 7 anchor mismatch. Expected 'Article Descriptors', got '{r1_col7}'")
    if r3_col8 != "Title":
        issues.append(f"Row 3, Col 8 anchor mismatch. Expected 'Title', got '{r3_col8}'")

    return len(issues) == 0, issues

def main():
    parser = argparse.ArgumentParser(description="Canonical 3-Tier Excel Coding Sheet Generator & Validator")
    parser.add_argument("--output", help="Create a blank 3-tier coding sheet at the specified path")
    parser.add_argument("--validate", help="Validate 3-tier header compliance of an existing Excel sheet")
    parser.add_argument("--test", action="store_true", help="Run self-test generating and validating a temp sheet")

    args = parser.parse_args()

    if args.test:
        test_path = os.path.join(os.path.dirname(DEFAULT_TEMPLATE_PATH), "_test_header_temp.xlsx")
        print(f"[TEST] Creating test sheet at {test_path}...")
        create_blank_coding_sheet(test_path)
        valid, issues = validate_coding_sheet_headers(test_path)
        if os.path.exists(test_path):
            os.remove(test_path)
        if valid:
            print("[TEST PASS] 3-tier header generation and validation verified successfully.")
            sys.exit(0)
        else:
            print(f"[TEST FAIL] Validation issues: {issues}")
            sys.exit(1)

    if args.output:
        create_blank_coding_sheet(args.output)
        print(f"Created canonical 3-tier coding sheet at: {args.output}")

    if args.validate:
        valid, issues = validate_coding_sheet_headers(args.validate)
        if valid:
            print(f"[PASS] '{args.validate}' conforms to canonical 3-tier header structure.")
        else:
            print(f"[FAIL] '{args.validate}' has {len(issues)} issue(s):")
            for iss in issues:
                print(f"  - {iss}")
            sys.exit(1)

if __name__ == "__main__":
    main()
